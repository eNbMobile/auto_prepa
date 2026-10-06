#!/usr/bin/env python3
"""
Crée le classeur "CA DRIVE <année>" à partir de celui de l'année précédente
(mêmes onglets, mêmes formules, mêmes paramètres), vidé de ses données réelles :

- EFFECTIF : seuls les salariés encore présents la dernière semaine de l'année
  précédente sont gardés, prolongés toute l'année à leur contrat ; mois
  recalés sur les semaines ISO de la nouvelle année.
- PRÉVI FCT CA MAG : tout vide sauf « Quota N-1 », qui reprend le quota réalisé
  de l'année précédente (formule sur RÉALISATION N-1).
- PRÉVI FCT N-1 : prévisions recalculées sur RÉALISATION N-1 (= le réalisé de
  l'année précédente), mêmes taux de progression ; jours fériés de la nouvelle
  année à 0, jours fériés de l'année précédente remplacés par la semaine d'avant.
- RÉALISATION N-1 : le réalisé de l'année précédente (CA, commandes, heures,
  produits, CA magasin).
- RÉALISATION, BESOIN HEURES, VHT, DIFF PRÉVIRÉAL : données vidées, formules
  conservées.
- ÉVOLUTION : SOMME recalées jour par jour sur les mois de la nouvelle année
  (et de l'année précédente pour les colonnes N-1).
- Cases noires : jours fériés de la nouvelle année.

Usage :
  creer_ca_drive_annee.py [--annee AAAA] [--remplacer]
      Télécharge "CA DRIVE <année-1>.xlsx" sur Drive, crée "CA DRIVE <année>.xlsx"
      dans le même dossier (refuse d'écraser un fichier existant sans --remplacer).
  creer_ca_drive_annee.py --annee AAAA --source SRC.xlsx --sortie DEST.xlsx
      Même transformation, en local (sans Drive).
"""
import io
import sys
from collections import Counter
from copy import copy
from datetime import date, datetime, timedelta

import openpyxl
from openpyxl.styles import PatternFill
from openpyxl.styles.cell_style import StyleArray
from openpyxl.utils import get_column_letter as L, column_index_from_string as I
from openpyxl.formula.translate import Translator
from openpyxl.worksheet.formula import ArrayFormula

MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

EFFECTIF = "EFFECTIF"
CA_MAG = "PRÉVI FCT CA MAG"
PREVI_N1 = "PRÉVI FCT N-1"
REAL_N1 = "RÉALISATION N-1"
REAL = "RÉALISATION"
DIFF = "DIFF PRÉVIRÉAL"
BESOIN = "BESOIN HEURES"
VHT = "VHT"
EVOL = "ÉVOLUTION"

NOIR = PatternFill("solid", fgColor="FF000000")
MOIS = ["JANVIER", "FÉVRIER", "MARS", "AVRIL", "MAI", "JUIN", "JUILLET",
        "AOÛT", "SEPTEMBRE", "OCTOBRE", "NOVEMBRE", "DÉCEMBRE"]

# Lignes des semaines (S 1 → S 53) : ligne = décalage + numéro de semaine.
LIGNE_S1 = 3          # RÉALISATION, RÉALISATION N-1, PRÉVI FCT N-1, DIFF, BESOIN, VHT
LIGNE_S1_MAG = 60     # bloc CA MAGASIN de RÉALISATION / RÉALISATION N-1
LIGNE_S1_CA_MAG = 4   # PRÉVI FCT CA MAG
NB_LIGNES_SEMAINES = 53
LIGNE_TOTAL = LIGNE_S1 + NB_LIGNES_SEMAINES          # 56
LIGNE_TOTAL_MAG = LIGNE_S1_MAG + NB_LIGNES_SEMAINES  # 113

# Colonnes du lundi au samedi de chaque bloc journalier.
def _jours(premiere):
    return [L(I(premiere) + k) for k in range(6)]

REAL_BLOCS = {"ca": _jours("B"), "cdes": _jours("L"), "pdts": _jours("U"),
              "ratio": _jours("AD"), "panier": _jours("AM")}
N1_BLOCS = {"ca": _jours("B"), "cdes": _jours("K"), "heures": _jours("S"),
            "pdts": _jours("AB"), "ratio": _jours("AK"), "panier": _jours("AT")}
BESOIN_BLOCS = {"prevu": _jours("B"), "planif": _jours("L"), "real": _jours("U"),
                "diff": _jours("AE")}
VHT_BLOCS = {"vht": _jours("B"), "quota": _jours("K"), "prod": _jours("T"),
             "possible": _jours("AC")}
CA_MAG_BLOCS = {"mag": _jours("B"), "drive": _jours("J"), "cdes": _jours("U")}


# ─────────────────────────────────────────────────────────────────
# Calendrier
# ─────────────────────────────────────────────────────────────────

def _paques(annee):
    a, b, c = annee % 19, annee // 100, annee % 100
    d, e = b // 4, b % 4
    f = (b + 8) // 25
    g = (b - f + 1) // 3
    h = (19 * a + b - d - g + 15) % 30
    i, k = c // 4, c % 4
    l = (32 + 2 * e + 2 * i - h - k) % 7
    m = (a + 11 * h + 22 * l) // 451
    mois = (h + l - 7 * m + 114) // 31
    jour = (h + l - 7 * m + 114) % 31 + 1
    return date(annee, mois, jour)


def jours_feries(annee):
    p = _paques(annee)
    return {
        date(annee, 1, 1): "Jour de l'an",
        p + timedelta(days=1): "Lundi de Pâques",
        date(annee, 5, 1): "Fête du travail",
        date(annee, 5, 8): "Victoire 1945",
        p + timedelta(days=39): "Ascension",
        p + timedelta(days=50): "Lundi de Pentecôte",
        date(annee, 7, 14): "Fête nationale",
        date(annee, 8, 15): "Assomption",
        date(annee, 11, 1): "Toussaint",
        date(annee, 11, 11): "Armistice",
        date(annee, 12, 25): "Noël",
    }


def drive_ferme(jour):
    """Le drive ferme les jours fériés, sauf le 8 mai (ouvert en 2026)."""
    return jour in jours_feries(jour.year) and (jour.month, jour.day) != (5, 8)


def magasin_ferme(jour):
    """Fériés où le magasin est fermé (vides en 2026 dans le bloc CA MAGASIN,
    plus Noël)."""
    if (jour.month, jour.day) in {(1, 1), (5, 1), (7, 14), (12, 25)}:
        return True
    return jour == _paques(jour.year) + timedelta(days=1)


def nb_semaines_iso(annee):
    return date(annee, 12, 28).isocalendar()[1]


def jours_ouvres_iso(annee):
    """(semaine, index du jour 0=lundi…5=samedi, date) de l'année ISO."""
    for s in range(1, nb_semaines_iso(annee) + 1):
        for j in range(6):
            yield s, j, date.fromisocalendar(annee, s, j + 1)


def mois_semaine(annee, semaine):
    """Mois (1-12) d'une semaine ISO : celui de son jeudi."""
    return date.fromisocalendar(annee, semaine, 4).month


# ─────────────────────────────────────────────────────────────────
# Outils feuilles
# ─────────────────────────────────────────────────────────────────

def _ref(onglet, cellule):
    return f"'{onglet}'!{cellule}"


def normaliser_styles(ws, lignes, colonnes):
    """Applique à chaque colonne le style le plus fréquent de la zone : efface
    les cases noires et les surlignages propres aux données de l'an passé."""
    for c in colonnes:
        styles = Counter(_style(ws.cell(r, c)) for r in lignes)
        mode = StyleArray(list(styles.most_common(1)[0][0]))
        for r in lignes:
            cell = ws.cell(r, c)
            if _style(cell) != tuple(mode):
                cell._style = copy(mode)


def _style(cell):
    return tuple(cell._style) if cell._style is not None else tuple(StyleArray())


def noircir(ws, cellule, vider=True):
    ws[cellule].fill = NOIR
    if vider:
        ws[cellule].value = None


def supprimer_commentaires(wb):
    for ws in wb:
        for row in ws.iter_rows():
            for cell in row:
                if cell.comment:
                    cell.comment = None


def remplacer_annee(ws, ancienne, nouvelle, lignes=(1,)):
    for r in lignes:
        for cell in ws[r]:
            if isinstance(cell.value, str) and str(ancienne) in cell.value:
                cell.value = cell.value.replace(str(ancienne), str(nouvelle))


def sans_div0(ws, plage):
    """Entoure de SIERREUR les formules de la plage (totaux et évolutions qui
    divisent par des cellules encore vides en début d'année)."""
    for ligne in ws[plage]:
        for cell in ligne:
            v = cell.value
            if isinstance(v, str) and v.startswith("=") and "/" in v and "IFERROR" not in v:
                cell.value = f"=IFERROR({v[1:]},0)"


def somme_jours(cellules):
    """Formule SOMME compacte d'une liste de (onglet, ligne, colonne)."""
    par_onglet = {}
    for onglet, ligne, col in cellules:
        par_onglet.setdefault(onglet, {}).setdefault(ligne, []).append(I(col))
    plages = []
    for onglet, lignes in par_onglet.items():
        segments = []  # (ligne_debut, ligne_fin, col_debut, col_fin)
        for ligne in sorted(lignes):
            cols = sorted(lignes[ligne])
            debut = prec = cols[0]
            for c in cols[1:] + [None]:
                if c is not None and c == prec + 1:
                    prec = c
                    continue
                seg = (ligne, ligne, debut, prec)
                if segments and segments[-1][1] == ligne - 1 and segments[-1][2:] == seg[2:]:
                    segments[-1] = (segments[-1][0], ligne, debut, prec)
                else:
                    segments.append(seg)
                if c is not None:
                    debut = prec = c
        for l1, l2, c1, c2 in segments:
            a, b = f"{L(c1)}{l1}", f"{L(c2)}{l2}"
            plages.append(_ref(onglet, a if a == b else f"{a}:{b}"))
    return f"=SUM({','.join(plages)})"


# ─────────────────────────────────────────────────────────────────
# Onglets
# ─────────────────────────────────────────────────────────────────

def maj_effectif(ws, annee):
    nb_sem = nb_semaines_iso(annee)
    col_s1 = 3
    derniere = col_s1 + NB_LIGNES_SEMAINES - 1  # colonne de la semaine 53 (BC)
    lignes_salaries = range(3, 27)

    # Salariés encore présents la dernière semaine de l'année précédente.
    nb_sem_prec = nb_semaines_iso(annee - 1)
    col_fin = col_s1 + nb_sem_prec - 1
    gardes_cdi, gardes_autres = [], []
    for r in lignes_salaries:
        nom = ws.cell(r, 1).value
        if not nom or ws.cell(r, col_fin).value in (None, ""):
            continue
        contrat = ws.cell(r, 2).value
        heures = contrat if isinstance(contrat, (int, float)) else None
        if heures is None:
            for c in range(col_fin, col_s1 - 1, -1):
                v = ws.cell(r, c).value
                if isinstance(v, (int, float)) and v > 0:
                    heures = v
                    break
        (gardes_autres if r >= 19 else gardes_cdi).append((nom, contrat, heures))

    for r in lignes_salaries:
        for c in range(1, derniere + 1):
            ws.cell(r, c).value = None
    normaliser_styles(ws, lignes_salaries, range(col_s1, derniere + 1))

    for debut, gardes in ((3, gardes_cdi), (19, gardes_autres)):
        for k, (nom, contrat, heures) in enumerate(gardes):
            r = debut + k
            ws.cell(r, 1).value = nom
            ws.cell(r, 2).value = contrat
            for s in range(nb_sem):
                ws.cell(r, col_s1 + s).value = heures

    # Semaines et mois de la nouvelle année.
    for s in range(1, NB_LIGNES_SEMAINES + 1):
        ws.cell(2, col_s1 + s - 1).value = s if s <= nb_sem else None
    for m in list(ws.merged_cells.ranges):
        if m.min_row == 1:
            ws.unmerge_cells(str(m))
    for c in range(col_s1, derniere + 1):
        ws.cell(1, c).value = None
    for mois in range(1, 13):
        sems = [s for s in range(1, nb_sem + 1) if mois_semaine(annee, s) == mois]
        c1, c2 = col_s1 + sems[0] - 1, col_s1 + sems[-1] - 1
        ws.cell(1, c1).value = MOIS[mois - 1]
        ws.merge_cells(start_row=1, start_column=c1, end_row=1, end_column=c2)

    for c in range(col_s1, derniere + 1):
        ws.cell(28, c).value = f"=SUM({L(c)}3:{L(c)}26)"

    # Formules matricielles : seule la cellule d'ancrage garde la formule ;
    # les valeurs « déversées » en cache bloqueraient l'expansion.
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, ArrayFormula):
                for ligne in ws[cell.value.ref]:
                    for autre in ligne:
                        if autre.coordinate != cell.coordinate:
                            autre.value = None


def maj_real_n1(wb, wbv, annee):
    """RÉALISATION N-1 = le réalisé de l'année précédente, 53 lignes de semaines
    (mêmes lignes que RÉALISATION), total, poids, puis le bloc CA MAGASIN."""
    ws, src, srcv, besoinv = wb[REAL_N1], wb[REAL], wbv[REAL], wbv[BESOIN]
    prec = annee - 1

    def ligne_libelle(texte):
        for r in range(3, ws.max_row + 1):
            if str(ws.cell(r, 1).value or "").strip().upper() == texte.upper():
                return r
        raise ValueError(f"{REAL_N1} : ligne {texte!r} introuvable")

    nb_col = I("AZ")
    st_semaine = [copy(ws.cell(LIGNE_S1 + 1, c)._style) for c in range(1, nb_col + 1)]
    st_total = [copy(ws.cell(ligne_libelle("Total"), c)._style) for c in range(1, nb_col + 1)]
    st_poids = [copy(ws.cell(ligne_libelle("Poids"), c)._style) for c in range(1, nb_col + 1)]
    st_entete_jours = copy(ws["B2"]._style)

    for m in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(m))
    for row in ws.iter_rows(min_row=1, max_row=max(ws.max_row, LIGNE_TOTAL_MAG)):
        for cell in row:
            cell.value = None
            if cell.row >= LIGNE_S1:
                cell._style = copy(ws.cell(1, 100)._style)

    titres = {"A": f"CA DRIVE {prec}", "J": f"NB COMMANDES DRIVE {prec}",
              "S": f"HEURES DRIVE RÉEL {prec}", "AA": f"NB PRODUITS DRIVE {prec}",
              "AJ": f"NB PRODUITS/CDE DRIVE {prec}", "AS": f"PANIER MOYEN DRIVE {prec}"}
    jours = ["Lundi", "Mardi", "Mercredi", "Jeudi ", "Vendredi ", "Samedi ", "Total"]
    for col, titre in titres.items():
        c = I(col)
        ws.cell(1, c).value = titre
        ws.cell(1, c)._style = copy(src["A1"]._style)
        largeur = 6 if col == "S" else 7
        ws.merge_cells(start_row=1, start_column=c, end_row=1, end_column=c + largeur)
        premier_jour = c if col == "S" else c + 1
        for k, j in enumerate(jours):
            ws.cell(2, premier_jour + k).value = j
            ws.cell(2, premier_jour + k)._style = copy(st_entete_jours)

    for s in range(1, NB_LIGNES_SEMAINES + 1):
        r = LIGNE_S1 + s - 1
        for c in range(1, nb_col + 1):
            ws.cell(r, c)._style = copy(st_semaine[c - 1])
        for col in ("A", "J", "AA", "AJ", "AS"):
            ws[f"{col}{r}"] = f"S {s}"
        for k in range(6):
            for bloc_n1, bloc_src, feuille in (("ca", REAL_BLOCS["ca"], srcv),
                                               ("cdes", REAL_BLOCS["cdes"], srcv),
                                               ("pdts", REAL_BLOCS["pdts"], srcv),
                                               ("heures", BESOIN_BLOCS["real"], besoinv)):
                v = feuille[f"{bloc_src[k]}{r}"].value
                ws[f"{N1_BLOCS[bloc_n1][k]}{r}"] = v if isinstance(v, (int, float)) else None
            ca, cde, pdt = N1_BLOCS["ca"][k], N1_BLOCS["cdes"][k], N1_BLOCS["pdts"][k]
            ws[f"{N1_BLOCS['ratio'][k]}{r}"] = f"=IF({cde}{r}<>0,{pdt}{r}/{cde}{r},0)"
            ws[f"{N1_BLOCS['panier'][k]}{r}"] = f"=IF({cde}{r}<>0,{ca}{r}/{cde}{r},0)"
        ws[f"H{r}"] = f"=SUM(B{r}:G{r})"
        ws[f"Q{r}"] = f"=SUM(K{r}:P{r})"
        ws[f"Y{r}"] = f"=SUM(S{r}:X{r})"
        ws[f"AH{r}"] = f"=SUM(AB{r}:AG{r})"
        ws[f"AQ{r}"] = f"=IF(Q{r}<>0,AH{r}/Q{r},0)"
        ws[f"AZ{r}"] = f"=IF(Q{r}<>0,H{r}/Q{r},0)"

    t, der = LIGNE_TOTAL, LIGNE_TOTAL - 1
    for c in range(1, nb_col + 1):
        ws.cell(t, c)._style = copy(st_total[c - 1])
        ws.cell(t + 1, c)._style = copy(st_poids[c - 1])
    for col in ("A", "J", "AA", "AJ", "AS"):
        ws[f"{col}{t}"] = "Total"
    for col in N1_BLOCS["ca"] + ["H"] + N1_BLOCS["cdes"] + ["Q"] + N1_BLOCS["heures"] + ["Y"] \
            + N1_BLOCS["pdts"] + ["AH"]:
        ws[f"{col}{t}"] = f"=SUM({col}{LIGNE_S1}:{col}{der})"
    for k in range(6):
        ca, cde, pdt = N1_BLOCS["ca"][k], N1_BLOCS["cdes"][k], N1_BLOCS["pdts"][k]
        ws[f"{N1_BLOCS['ratio'][k]}{t}"] = f"=IF({cde}{t}<>0,{pdt}{t}/{cde}{t},0)"
        ws[f"{N1_BLOCS['panier'][k]}{t}"] = f"=IF({cde}{t}<>0,{ca}{t}/{cde}{t},0)"
        ws[f"{ca}{t + 1}"] = f"=IF(H{t}<>0,{ca}{t}/H{t},0)"
        ws[f"{cde}{t + 1}"] = f"=IF(Q{t}<>0,{cde}{t}/Q{t},0)"
    ws[f"AQ{t}"] = f"=IF(Q{t}<>0,AH{t}/Q{t},0)"
    ws[f"AZ{t}"] = f"=IF(Q{t}<>0,H{t}/Q{t},0)"
    ws[f"A{t + 1}"] = "Poids"
    ws[f"J{t + 1}"] = "Poids"

    # Bloc CA MAGASIN, même disposition que dans RÉALISATION.
    titre_mag = LIGNE_S1_MAG - 2
    for r in range(titre_mag, LIGNE_TOTAL_MAG + 1):
        for c in range(1, 9):
            ws.cell(r, c)._style = copy(src.cell(r, c)._style)
            v = src.cell(r, c).value
            if r < LIGNE_S1_MAG or r == LIGNE_TOTAL_MAG or c in (1, 8):
                ws.cell(r, c).value = v
            else:
                v = srcv.cell(r, c).value
                ws.cell(r, c).value = v if isinstance(v, (int, float)) else None
    ws.cell(titre_mag, 1).value = f"CA MAGASIN {prec}"
    normaliser_styles(ws, range(LIGNE_S1_MAG, LIGNE_S1_MAG + NB_LIGNES_SEMAINES), range(1, 9))
    ws.merge_cells(start_row=titre_mag, start_column=1, end_row=titre_mag, end_column=8)


def maj_previ_ca_mag(ws, annee):
    remplacer_annee(ws, annee - 1, annee)
    lignes = range(LIGNE_S1_CA_MAG, LIGNE_S1_CA_MAG + NB_LIGNES_SEMAINES)
    normaliser_styles(ws, lignes, range(1, I("AC") + 1))
    for s in range(1, NB_LIGNES_SEMAINES + 1):
        r = LIGNE_S1_CA_MAG + s - 1
        for col in CA_MAG_BLOCS["mag"]:
            ws[f"{col}{r}"] = None
        ws[f"R{r}"] = None
        n1, mag = LIGNE_S1 + s - 1, LIGNE_S1_MAG + s - 1
        ws[f"Q{r}"] = (f"=IF({_ref(REAL_N1, f'H{mag}')}<>0,"
                       f"{_ref(REAL_N1, f'H{n1}')}/{_ref(REAL_N1, f'H{mag}')},\"\")")
    for s, k, jour in jours_ouvres_iso(annee):
        r = LIGNE_S1_CA_MAG + s - 1
        if drive_ferme(jour):
            noircir(ws, f"{CA_MAG_BLOCS['drive'][k]}{r}", vider=False)
            noircir(ws, f"{CA_MAG_BLOCS['cdes'][k]}{r}", vider=False)
        if magasin_ferme(jour):
            noircir(ws, f"{CA_MAG_BLOCS['mag'][k]}{r}")


def maj_previ_n1(ws, annee):
    remplacer_annee(ws, annee - 1, annee)
    lignes = range(LIGNE_S1, LIGNE_S1 + NB_LIGNES_SEMAINES)
    normaliser_styles(ws, lignes, range(1, I("AS") + 1))
    nb_sem, nb_sem_prec = nb_semaines_iso(annee), nb_semaines_iso(annee - 1)

    def ligne_source(s, k):
        """Ligne de RÉALISATION N-1 servant de base : la même semaine, ou la
        précédente (suivante en S 1) si ce jour-là était férié l'an passé."""
        pas = -1 if s > 1 else 1
        while 1 <= s <= nb_sem_prec and drive_ferme(date.fromisocalendar(annee - 1, s, k + 1)):
            s += pas
        return LIGNE_S1 + s - 1

    for s in range(1, NB_LIGNES_SEMAINES + 1):
        r = LIGNE_S1 + s - 1
        for k in range(6):
            cellules = [REAL_BLOCS[b][k] for b in ("ca", "cdes", "pdts", "ratio", "panier")]
            if s > nb_sem:
                for col in cellules[:3]:
                    ws[f"{col}{r}"] = None
            elif drive_ferme(date.fromisocalendar(annee, s, k + 1)):
                for col in cellules:
                    ws[f"{col}{r}"] = 0
                    noircir(ws, f"{col}{r}", vider=False)
                continue
            else:
                src = ligne_source(s, k)
                ws[f"{REAL_BLOCS['ca'][k]}{r}"] = \
                    f"={_ref(REAL_N1, N1_BLOCS['ca'][k] + str(src))}*(1+J{r}+1%)"
                ws[f"{REAL_BLOCS['cdes'][k]}{r}"] = \
                    f"={_ref(REAL_N1, N1_BLOCS['cdes'][k] + str(src))}*(1+J{r})"
                ws[f"{REAL_BLOCS['pdts'][k]}{r}"] = \
                    f"={_ref(REAL_N1, N1_BLOCS['pdts'][k] + str(src))}*(1+J{r}+1%)"
            ca, cde, pdt = (REAL_BLOCS[b][k] for b in ("ca", "cdes", "pdts"))
            ws[f"{REAL_BLOCS['ratio'][k]}{r}"] = f"=IF({cde}{r}<>0,{pdt}{r}/{cde}{r},0)"
            ws[f"{REAL_BLOCS['panier'][k]}{r}"] = f"=IF({cde}{r}<>0,{ca}{r}/{cde}{r},0)"
        ws[f"AJ{r}"] = f"=IF(R{r}<>0,AA{r}/R{r},0)"
        ws[f"AS{r}"] = f"=IF(R{r}<>0,H{r}/R{r},0)"
        if s > nb_sem:
            ws[f"J{r}"] = None

    # Totaux commandes calculés sur le total de RÉALISATION N-1 (ligne Total).
    t = LIGNE_TOTAL
    for k, col in enumerate(REAL_BLOCS["cdes"]):
        ws[f"{col}{t}"] = f"={_ref(REAL_N1, N1_BLOCS['cdes'][k] + str(t))}*(1+J{t})"


def maj_realisation(ws, annee):
    lignes = list(range(LIGNE_S1, LIGNE_S1 + NB_LIGNES_SEMAINES))
    lignes_mag = list(range(LIGNE_S1_MAG, LIGNE_S1_MAG + NB_LIGNES_SEMAINES))
    normaliser_styles(ws, lignes, range(1, I("AS") + 1))
    normaliser_styles(ws, lignes_mag, range(1, I("H") + 1))
    for r in lignes:
        for k in range(6):
            for b in ("ca", "cdes", "pdts"):
                ws[f"{REAL_BLOCS[b][k]}{r}"] = None
            ca, cde, pdt = (REAL_BLOCS[b][k] for b in ("ca", "cdes", "pdts"))
            ws[f"{REAL_BLOCS['ratio'][k]}{r}"] = f"=IF({cde}{r}<>0,{pdt}{r}/{cde}{r},0)"
            ws[f"{REAL_BLOCS['panier'][k]}{r}"] = f"=IF({cde}{r}<>0,{ca}{r}/{cde}{r},0)"
        ws[f"AJ{r}"] = f"=IF(R{r}<>0,AA{r}/R{r},0)"
        ws[f"AS{r}"] = f"=IF(R{r}<>0,H{r}/R{r},0)"
    for r in lignes_mag:
        for col in REAL_BLOCS["ca"]:
            ws[f"{col}{r}"] = None
    for s, k, jour in jours_ouvres_iso(annee):
        r = LIGNE_S1 + s - 1
        if drive_ferme(jour):
            for b in REAL_BLOCS:
                noircir(ws, f"{REAL_BLOCS[b][k]}{r}")
        if magasin_ferme(jour):
            noircir(ws, f"{REAL_BLOCS['ca'][k]}{LIGNE_S1_MAG + s - 1}")


def maj_diff(ws, annee):
    remplacer_annee(ws, annee - 1, annee)
    lignes = range(LIGNE_S1, LIGNE_S1 + nb_semaines_iso(annee - 1))
    lignes = [r for r in lignes if str(ws.cell(r, 1).value or "").startswith("S")]
    normaliser_styles(ws, lignes, range(1, I("AS") + 1))
    for r in lignes:
        for b in REAL_BLOCS:
            for col in REAL_BLOCS[b]:
                re_, pr = _ref(REAL, f"{col}{r}"), _ref(PREVI_N1, f"{col}{r}")
                ws[f"{col}{r}"] = f"=IF({re_}<>0,{re_}-{pr},0)"
        for col in ("AJ", "AS"):
            re_, pr = _ref(REAL, f"{col}{r}"), _ref(PREVI_N1, f"{col}{r}")
            ws[f"{col}{r}"] = f"=IF({re_}<>0,{re_}-{pr},0)"
    for s, k, jour in jours_ouvres_iso(annee):
        r = LIGNE_S1 + s - 1
        if r in lignes and drive_ferme(jour):
            for b in REAL_BLOCS:
                noircir(ws, f"{REAL_BLOCS[b][k]}{r}")


def maj_besoin(ws, annee):
    lignes = range(LIGNE_S1, LIGNE_S1 + NB_LIGNES_SEMAINES)
    normaliser_styles(ws, lignes, range(1, I("AK") + 1))
    for r in lignes:
        for k in range(6):
            mag = _ref(CA_MAG, f"{CA_MAG_BLOCS['drive'][k]}{r + 1}")
            n1 = _ref(PREVI_N1, f"{REAL_BLOCS['ca'][k]}{r}")
            ws[f"{BESOIN_BLOCS['prevu'][k]}{r}"] = \
                f"=IF({mag}<>0,({n1}*0.75+{mag}*0.25)/J{r},0)"
            for b in ("planif", "real"):
                ws[f"{BESOIN_BLOCS[b][k]}{r}"] = None
            ws[f"{BESOIN_BLOCS['diff'][k]}{r}"] = \
                f"={BESOIN_BLOCS['real'][k]}{r}-{BESOIN_BLOCS['planif'][k]}{r}"
        ws[f"R{r}"] = f"=SUM(L{r}:Q{r})"
        ws[f"AA{r}"] = f"=SUM(U{r}:Z{r})"
        ws[f"AK{r}"] = f"=SUM(AE{r}:AJ{r})"
    for s, k, jour in jours_ouvres_iso(annee):
        r = LIGNE_S1 + s - 1
        if drive_ferme(jour):
            for b in BESOIN_BLOCS:
                noircir(ws, f"{BESOIN_BLOCS[b][k]}{r}")


def maj_vht(ws, annee):
    lignes = range(LIGNE_S1, LIGNE_S1 + NB_LIGNES_SEMAINES)
    normaliser_styles(ws, lignes, range(1, I("AI") + 1))
    for r in lignes:
        m = r + LIGNE_S1_MAG - LIGNE_S1
        for k in range(7):
            ca = L(I("B") + k)
            heures = _ref(BESOIN, f"{L(I('U') + k)}{r}")
            ws[f"{L(I('B') + k)}{r}"] = \
                f"=IF({heures}<>0,{_ref(REAL, f'{ca}{r}')}/({heures}*0.9523809524),0)"
            ws[f"{L(I('K') + k)}{r}"] = (f"=IF({_ref(REAL, f'{ca}{r}')}<>0,"
                                        f"{_ref(REAL, f'{ca}{r}')}/{_ref(REAL, f'{ca}{m}')},0)")
            pdts = _ref(REAL, f"{L(I('U') + k)}{r}")
            ws[f"{L(I('T') + k)}{r}"] = f"=IF({pdts}<>0,{pdts}/{heures},0)"
            if k < 6:
                ratio = _ref(PREVI_N1, f"{REAL_BLOCS['ratio'][k]}{r}")
                planif = _ref(BESOIN, f"{BESOIN_BLOCS['planif'][k]}{r}")
                ws[f"{VHT_BLOCS['possible'][k]}{r}"] = f"=IF({ratio}<>0,{planif}*75/{ratio},0)"
        ws[f"AI{r}"] = f"=SUM(AC{r}:AH{r})"
    for s, k, jour in jours_ouvres_iso(annee):
        r = LIGNE_S1 + s - 1
        if drive_ferme(jour):
            for b in VHT_BLOCS:
                noircir(ws, f"{VHT_BLOCS[b][k]}{r}")


def maj_evolution(ws, annee):
    remplacer_annee(ws, annee - 1, annee)
    prec = annee - 1
    ws["B19"], ws["E19"] = prec, annee

    def cellules_mois(a, mois, bloc_real, bloc_n1, mag=False):
        """Cellules des jours (lundi → samedi) du mois `mois` de l'année `a`,
        dans RÉALISATION si la semaine ISO est de l'année du classeur, sinon
        dans RÉALISATION N-1 (semaine ISO de l'année précédente)."""
        cellules = []
        jour = date(a, mois, 1)
        while jour.month == mois:
            annee_iso, s, j = jour.isocalendar()
            if j <= 6:
                ligne = (LIGNE_S1_MAG if mag else LIGNE_S1) + s - 1
                if annee_iso == annee:
                    cellules.append((REAL, ligne, bloc_real[j - 1]))
                elif annee_iso == prec:
                    cellules.append((REAL_N1, ligne, bloc_n1[j - 1]))
            jour += timedelta(days=1)
        return cellules

    for mois in range(1, 13):
        r = 2 + mois
        for col, a, bloc in (("B", annee, "ca"), ("C", prec, "ca"),
                             ("E", annee, "cdes"), ("F", prec, "cdes"),
                             ("N", annee, "pdts"), ("O", prec, "pdts")):
            ws[f"{col}{r}"] = somme_jours(cellules_mois(a, mois, REAL_BLOCS[bloc], N1_BLOCS[bloc]))
        rm = 19 + mois
        ws[f"B{rm}"] = somme_jours(cellules_mois(prec, mois, REAL_BLOCS["ca"], N1_BLOCS["ca"], mag=True))
        ws[f"E{rm}"] = somme_jours(cellules_mois(annee, mois, REAL_BLOCS["ca"], N1_BLOCS["ca"], mag=True))
    # Janvier : mêmes formules d'évolution que les autres mois.
    for col in ("D", "G", "J", "K", "M", "P"):
        ws[f"{col}3"] = Translator(ws[f"{col}4"].value, origin=f"{col}4").translate_formula(f"{col}3")


# ─────────────────────────────────────────────────────────────────

def transformer(contenu, annee):
    wb = openpyxl.load_workbook(io.BytesIO(contenu))
    wbv = openpyxl.load_workbook(io.BytesIO(contenu), data_only=True)
    supprimer_commentaires(wb)
    maj_effectif(wb[EFFECTIF], annee)
    maj_real_n1(wb, wbv, annee)          # avant de vider RÉALISATION
    maj_previ_ca_mag(wb[CA_MAG], annee)
    maj_previ_n1(wb[PREVI_N1], annee)
    maj_realisation(wb[REAL], annee)
    maj_diff(wb[DIFF], annee)
    maj_besoin(wb[BESOIN], annee)
    maj_vht(wb[VHT], annee)
    maj_evolution(wb[EVOL], annee)
    sans_div0(wb[REAL], f"AD{LIGNE_TOTAL}:AS{LIGNE_TOTAL}")
    sans_div0(wb[VHT], f"B{LIGNE_TOTAL}:AI{LIGNE_TOTAL}")
    sans_div0(wb[EVOL], "D3:P15")
    sans_div0(wb[EVOL], "H20:M31")
    wb.calculation.fullCalcOnLoad = True
    sortie = io.BytesIO()
    wb.save(sortie)
    return sortie.getvalue()


# ─────────────────────────────────────────────────────────────────
# Drive
# ─────────────────────────────────────────────────────────────────

def _drive_creer(annee, remplacer):
    import controle_stocks as cs
    import renseigne_ca as rc
    from googleapiclient.http import MediaIoBaseUpload

    cs._charger_config()
    drive = cs._get_drive_service()
    if not drive:
        print("ERREUR : Drive inaccessible.")
        sys.exit(1)
    source = rc.id_classeur_ca(drive, annee - 1)
    meta = drive.files().get(fileId=source, fields="name,parents",
                             supportsAllDrives=True).execute()
    if str(annee - 1) not in meta["name"]:
        print(f"ERREUR : le classeur source s'appelle {meta['name']!r}, "
              f"pas « CA DRIVE {annee - 1} ».")
        sys.exit(1)
    # Classeur partagé dont le dossier est invisible : racine du Drive.
    dossier = (meta.get("parents") or [None])[0]
    nom = f"CA DRIVE {annee}.xlsx"
    existant = rc.id_classeur_ca(drive, annee)
    existants = [{"id": existant}] if existant else []
    if existants and not remplacer:
        print(f"ERREUR : {nom} existe déjà sur Drive (relancer avec --remplacer).")
        sys.exit(1)

    print(f"Source : {meta['name']} → {nom}")
    contenu = transformer(rc._telecharger(drive, source), annee)
    media = MediaIoBaseUpload(io.BytesIO(contenu), mimetype=MIME_XLSX, resumable=False)
    if existants:
        f = drive.files().update(fileId=existants[0]["id"], media_body=media,
                                 supportsAllDrives=True).execute()
    else:
        corps = {"name": nom, "mimeType": MIME_XLSX}
        if dossier:
            corps["parents"] = [dossier]
        f = drive.files().create(body=corps, media_body=media, fields="id",
                                 supportsAllDrives=True).execute()
    print(f"{nom} {'remplacé' if existants else 'créé'} : "
          f"https://drive.google.com/file/d/{f['id']}/view")


def main():
    args = sys.argv[1:]

    def option(nom):
        if nom in args:
            i = args.index(nom)
            if i + 1 < len(args):
                return args[i + 1]
            print(f"Valeur manquante pour {nom}")
            sys.exit(1)
        return None

    annee = int(option("--annee") or datetime.now().year + 1)
    source, dest = option("--source"), option("--sortie")
    if source or dest:
        if not (source and dest):
            print("--source et --sortie vont ensemble.")
            sys.exit(1)
        with open(source, "rb") as f:
            contenu = transformer(f.read(), annee)
        with open(dest, "wb") as f:
            f.write(contenu)
        print(f"{dest} écrit (CA DRIVE {annee}).")
        return
    _drive_creer(annee, "--remplacer" in args)


if __name__ == "__main__":
    main()
